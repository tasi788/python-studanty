from operator import itemgetter
import openpyxl

wb = openpyxl.load_workbook('names_and_savings.xlsx')

ws = wb['工作表1']

results = []

for row in ws.iter_rows():
    results.append([cell.value for cell in row])

print(results)

print(ws.cell(row=1, column=1).value)
print(ws.cell(row=1, column=2).value)
print(ws.cell(row=1, column=3).value)
print(type(ws.cell(row=1, column=1).value))
print(type(ws.cell(row=1, column=2).value))
print(type(ws.cell(row=1, column=3).value))


ws.cell(row=1, column=4).value = 'VIP'
print(ws.cell(row=1, column=4).value)

wb.save('names_and_savings_new.xlsx')
